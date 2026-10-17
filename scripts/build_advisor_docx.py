from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/thesis/final/proposal-clean.md"
OUT = ROOT / "docs/thesis/final/Proposal_Skripsi_Faisal_Hidayat_Sukma_Advisor.docx"
ASSETS = ROOT / "docs/thesis/final/assets"

TITLE = "Rancang Bangun Autonomous Liquidity Provider Agent Berbasis Data On-Chain pada Jaringan EVM"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_table_borders(table, color='D9D9D9', size='6'):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f'w:{edge}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)

def add_page_field(paragraph, fmt=None):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld_sep = OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])

def add_page_number_footer(section, roman=False):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_page_field(p)
    if roman:
        # LibreOffice/Word field formatting is controlled by section page-number format.
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn('w:pgNumType'))
        if pg_num is None:
            pg_num = OxmlElement('w:pgNumType'); sect_pr.append(pg_num)
        pg_num.set(qn('w:fmt'), 'lowerRoman')
    else:
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn('w:pgNumType'))
        if pg_num is None:
            pg_num = OxmlElement('w:pgNumType'); sect_pr.append(pg_num)
        pg_num.set(qn('w:fmt'), 'decimal')
        pg_num.set(qn('w:start'), '1')

def continue_page_numbering(section):
    pg_num = section._sectPr.find(qn('w:pgNumType'))
    if pg_num is not None:
        section._sectPr.remove(pg_num)

def add_field(paragraph, code):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = f' {code} '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t'); txt.text = ''
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, sep, txt, end])

def style_doc(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size, bold, align in [('Heading 1', 14, True, WD_ALIGN_PARAGRAPH.CENTER), ('Heading 2', 12, True, WD_ALIGN_PARAGRAPH.LEFT), ('Heading 3', 12, True, WD_ALIGN_PARAGRAPH.LEFT)]:
        st = styles[name]; st.font.name = 'Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.font.size = Pt(size); st.font.bold = bold; st.font.color.rgb = None
        st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.alignment = align

def add_plain_paragraph(doc, text, align=None, indent=True, bold=False, size=12, italic=False):
    p = doc.add_paragraph(style='Normal')
    if align is not None: p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    if bold or italic or '**' in text:
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if not part: continue
            is_bold = part.startswith('**') and part.endswith('**')
            value = part[2:-2] if is_bold else part
            r = p.add_run(value); r.bold = is_bold or bold; r.italic = italic
            r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    else:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    return p

def add_title_page(doc, title, subtitle=None):
    for _ in range(7): doc.add_paragraph('')
    p = add_plain_paragraph(doc, title.upper(), WD_ALIGN_PARAGRAPH.CENTER, False, True, 14)
    p.paragraph_format.first_line_indent = Cm(0)
    if subtitle:
        add_plain_paragraph(doc, subtitle, WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    for _ in range(4): doc.add_paragraph('')
    for label, value in [('Nama', 'Faisal Hidayat Sukma'), ('NPM', '2310631250053'), ('Kelas', '7A Sistem Informasi')]:
        add_plain_paragraph(doc, f'{label}: {value}', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)

def add_front_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text.upper()); r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(14)

def add_toc_page(doc, title, field):
    add_front_heading(doc, title)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, field)

def add_static_index(doc, title, entries):
    add_front_heading(doc, title)
    for label, page in entries:
        p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(4)
        r = p.add_run(label); r.font.name = 'Times New Roman'; r.font.size = Pt(11)
        if page:
            tabs = p.paragraph_format.tab_stops
            tabs.add_tab_stop(Cm(15.0))
            r = p.add_run('\t' + page); r.font.name = 'Times New Roman'; r.font.size = Pt(11)

TABLE_LABELS = ['1.1', '2.1', '3.1', None, '3.2', '3.3', '3.4', '3.5', '3.6']
TABLE_TITLES = [
    'Jadwal Penelitian',
    'Ringkasan Penelitian Terdahulu',
    'Jaringan dan Protokol yang Didukung Viero',
    'Tahapan dan Rancangan Penelitian',
    'Komponen dan Tanggung Jawab Arsitektur Viero Agent',
    'Pipeline Keputusan dan Execution Viero Agent',
    'Lapisan Pemeriksaan Safety',
    'Level Pengujian dan Evaluasi Artefak',
    'Matriks Skenario dan Evaluasi E1–E12',
]

TABLE_WIDTHS = {
    1: [0.7, 3.2, 1.683, 1.683, 1.683, 1.683, 1.683, 1.683],
    2: [0.7, 2.8, 2.2, 2.4, 3.0, 3.1],
    3: [3.0, 1.5, 9.7],
    4: [3.4, 7.0, 3.8],
    5: [4.2, 10.0],
    6: [3.0, 4.3, 6.9],
    7: [3.0, 7.0, 4.2],
    8: [3.0, 3.5, 4.5, 3.2],
    9: [0.8, 3.8, 7.0, 3.7, 3.2, 1.1, 2.0],
}

def parse_table(lines, doc):
    rows = []
    for line in lines:
        if set(line.replace('|','').replace(':','').replace('-','').strip()) == set():
            continue
        cells = [clean_text(c.strip()) for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    if not rows: return
    cols = max(len(r) for r in rows)
    header_key = tuple(rows[0][:2])
    table_index = {
        ('No.', 'Kegiatan Penelitian'): 0,
        ('No.', 'Peneliti dan Tahun'): 1,
        ('Jaringan', 'Chain ID'): 2,
        ('Tahap', 'Kegiatan Utama'): 3,
        ('Komponen', 'Tanggung Jawab'): 4,
        ('Tahap', 'Input/Evidence'): 5,
        ('Tahap Safety', 'Contoh Pemeriksaan'): 6,
        ('Level', 'Fokus'): 7,
        ('ID', 'Skenario'): 8,
    }.get(header_key)
    label = TABLE_LABELS[table_index] if table_index is not None else None
    caption = TABLE_TITLES[table_index] if table_index is not None else 'Tabel Penelitian'
    landscape_section = None
    if table_index == 8:
        landscape_section = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape_section.orientation = WD_ORIENT.LANDSCAPE
        landscape_section.page_width = Cm(29.7); landscape_section.page_height = Cm(21)
        landscape_section.left_margin = Cm(4); landscape_section.right_margin = Cm(3)
        landscape_section.top_margin = Cm(3); landscape_section.bottom_margin = Cm(3)
        landscape_section.footer.is_linked_to_previous = True
        continue_page_numbering(landscape_section)
    if label:
        cp = add_plain_paragraph(doc, f'Tabel {label} {caption}', WD_ALIGN_PARAGRAPH.CENTER, False, True, 12)
        cp.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = TABLE_WIDTHS.get((table_index + 1) if table_index is not None else 0)
    if widths and len(widths) == cols:
        grid = table._tbl.tblGrid
        grid_cols = list(grid)
        for node in grid_cols:
            grid.remove(node)
        for width in widths:
            col = OxmlElement('w:gridCol')
            col.set(qn('w:w'), str(int(width * 567)))
            grid.append(col)
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn('w:tblW'))
        if tbl_w is None:
            tbl_w = OxmlElement('w:tblW'); tbl_pr.append(tbl_w)
        tbl_w.set(qn('w:w'), str(int(sum(widths) * 567)))
        tbl_w.set(qn('w:type'), 'dxa')
        for row in table.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Cm(width)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j in range(cols):
        cell = hdr.cells[j]; set_cell_shading(cell, 'D9E2F3')
        set_cell_margins(cell, start=35, end=35) if table_index == 0 else set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths and j < len(widths): cell.width = Cm(widths[j])
        val = rows[0][j] if j < len(rows[0]) else ''
        p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean_text(val)); r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    for i, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for j in range(cols):
            cell = cells[j]
            set_cell_margins(cell, start=35, end=35) if table_index == 0 else set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths and j < len(widths): cell.width = Cm(widths[j])
            if i % 2 == 1: set_cell_shading(cell, 'F7F9FC')
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Cm(0)
            val = row[j] if j < len(row) else ''
            r = p.add_run(clean_text(val)); r.font.name = 'Times New Roman'; r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    if landscape_section is not None:
        restore = doc.add_section(WD_SECTION.NEW_PAGE)
        restore.orientation = WD_ORIENT.PORTRAIT
        restore.page_width = Cm(21); restore.page_height = Cm(29.7)
        restore.left_margin = Cm(4); restore.right_margin = Cm(3)
        restore.top_margin = Cm(3); restore.bottom_margin = Cm(3)
        restore.footer.is_linked_to_previous = True
        continue_page_numbering(restore)

def clean_text(s):
    s = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', s)
    s = s.replace('`', '')
    s = re.sub(r'(?<!\*)\*(?!\*)', '', s)
    return s.strip()

def add_content(doc):
    parse_table.counter = 0
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    in_body = False; table_buf = []; pending_image = None; first_chapter = True
    figure_map = {'Gambar 3.1': 'gambar-3-1.png', 'Gambar 3.2': 'gambar-3-2.png', 'Gambar 3.3': 'gambar-3-3.png'}
    for raw in lines:
        line = raw.strip()
        if line == '# BAB I PENDAHULUAN': in_body = True
        if not in_body: continue
        if line == '# DAFTAR PUSTAKA':
            if table_buf: parse_table(table_buf, doc); table_buf=[]
            p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading('DAFTAR PUSTAKA', level=1)
            continue
        if line.startswith('|'):
            table_buf.append(line); continue
        if table_buf:
            parse_table(table_buf, doc); table_buf=[]
        if not line: continue
        if line.startswith('*Gambar ') and line.endswith('*'):
            cap = line.strip('*')
            for key, filename in figure_map.items():
                if cap.startswith(key):
                    pimg = doc.add_paragraph(); pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER; pimg.paragraph_format.first_line_indent = Cm(0)
                    pimg.add_run().add_picture(str(ASSETS / filename), width=Cm(14.5))
                    break
            p = add_plain_paragraph(doc, cap, WD_ALIGN_PARAGRAPH.CENTER, False, True, 12)
            p.paragraph_format.keep_with_next = False
            continue
        if line.startswith('# '):
            text = clean_text(line[2:])
            if text == 'Proposal Tugas Akhir': continue
            if text in ('BAB I PENDAHULUAN', 'BAB II LANDASAN TEORI', 'BAB III OBJEK DAN METODOLOGI PENELITIAN'):
                if not first_chapter: doc.add_page_break()
                doc.add_heading(text, level=1); first_chapter = False
            elif re.match(r'^[123]\.\d+', text):
                doc.add_heading(text, level=2)
            else:
                doc.add_heading(text, level=1)
            continue
        if line.startswith('## '):
            doc.add_heading(clean_text(line[3:]), level=3); continue
        if line.startswith('### '):
            doc.add_heading(clean_text(line[4:]), level=3); continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.5
            r = p.add_run(clean_text(line[2:])); r.font.name='Times New Roman'; r.font.size=Pt(12); continue
        if re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number'); p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.5
            r = p.add_run(clean_text(re.sub(r'^\d+\. ', '', line))); r.font.name='Times New Roman'; r.font.size=Pt(12); continue
        add_plain_paragraph(doc, clean_text(line))
    if table_buf: parse_table(table_buf, doc)

def add_bibliography_format(doc):
    start = False
    for p in doc.paragraphs:
        if p.text.strip() == 'DAFTAR PUSTAKA': start = True; continue
        if start and p.text.strip():
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs: r.font.size = Pt(11)

def build():
    doc = Document()
    style_doc(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(4); sec.right_margin = Cm(3); sec.top_margin = Cm(3); sec.bottom_margin = Cm(3)
    # Cover
    add_page_number_footer(sec, roman=True)
    sec.different_first_page_header_footer = True
    for _ in range(5): doc.add_paragraph('')
    add_plain_paragraph(doc, 'PROPOSAL TUGAS AKHIR', WD_ALIGN_PARAGRAPH.CENTER, False, True, 14)
    for _ in range(3): doc.add_paragraph('')
    add_plain_paragraph(doc, TITLE.upper(), WD_ALIGN_PARAGRAPH.CENTER, False, True, 14)
    for _ in range(4): doc.add_paragraph('')
    for text in ['Faisal Hidayat Sukma', '2310631250053', '7A Sistem Informasi', 'Program Studi Sistem Informasi', 'Fakultas Ilmu Komputer', 'Universitas Singaperbangsa Karawang', '2026/2027']:
        add_plain_paragraph(doc, text, WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    doc.add_page_break()
    # Title page
    add_title_page(doc, 'Proposal Tugas Akhir', 'Untuk memenuhi salah satu syarat penyusunan tugas akhir')
    add_plain_paragraph(doc, TITLE, WD_ALIGN_PARAGRAPH.CENTER, False, True, 13)
    for _ in range(5): doc.add_paragraph('')
    add_plain_paragraph(doc, 'Faisal Hidayat Sukma', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    add_plain_paragraph(doc, '2310631250053', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    doc.add_page_break()
    # Approval page
    add_front_heading(doc, 'Lembar Pengesahan')
    add_plain_paragraph(doc, 'Proposal tugas akhir ini telah diperiksa dan disetujui untuk dikonsultasikan kepada dosen pembimbing.', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    for _ in range(8): doc.add_paragraph('')
    add_plain_paragraph(doc, 'Pembimbing,', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    for _ in range(4): doc.add_paragraph('')
    add_plain_paragraph(doc, 'Dr. Ade Andri Hendriadi, S.Si., M.Kom.', WD_ALIGN_PARAGRAPH.CENTER, False, False, 12)
    doc.add_page_break()
    # Preface
    add_front_heading(doc, 'Kata Pengantar')
    add_plain_paragraph(doc, 'Puji syukur penulis panjatkan kepada Tuhan Yang Maha Esa karena proposal tugas akhir ini dapat disusun sebagai dasar konsultasi penelitian. Proposal ini membahas perancangan dan evaluasi Viero Agent sebagai Autonomous Liquidity Provider Agent berbasis data on-chain pada jaringan EVM.', indent=True)
    add_plain_paragraph(doc, 'Penulis menyadari bahwa proposal ini masih dapat disempurnakan melalui arahan dosen pembimbing. Oleh karena itu, kritik dan saran yang membangun sangat diharapkan untuk memperbaiki rancangan penelitian dan penyajiannya.', indent=True)
    add_plain_paragraph(doc, 'Karawang, 2026', WD_ALIGN_PARAGRAPH.RIGHT, False, False, 12)
    add_plain_paragraph(doc, 'Penulis', WD_ALIGN_PARAGRAPH.RIGHT, False, False, 12)
    doc.add_page_break()
    toc_entries = [
        ('BAB I PENDAHULUAN', '1'), ('1.1 Latar Belakang', '1'), ('1.2 Rumusan Masalah', '3'), ('1.3 Batasan Masalah', '4'),
        ('1.4 Tujuan Penelitian', '6'), ('1.5 Manfaat Penelitian', '6'), ('1.5.1 Manfaat Teoretis', '6'), ('1.5.2 Manfaat Praktis', '7'),
        ('1.6 Metodologi Penelitian', '7'), ('1.7 Sistematika Penulisan', '8'), ('1.8 Jadwal Penelitian', '9'),
        ('BAB II LANDASAN TEORI', '11'), ('2.1 Blockchain, Jaringan EVM, dan Smart Contract', '11'), ('2.2 DeFi, DEX, AMM, Liquidity Pool, dan Liquidity Provider', '13'),
        ('2.3 Concentrated Liquidity dan Protokol yang Didukung Viero', '16'), ('2.4 Data On-Chain, RPC, dan Verifikasi State', '19'), ('2.5 Autonomous Software Agent dan Deterministic/Rule-Based Decision', '21'),
        ('2.6 Transaction Execution, Monitoring, Safety Boundary, dan Signer Isolation', '24'), ('2.7 Persistence, Reconciliation, dan Telegram sebagai Operator Interface', '26'),
        ('2.8 Penelitian Terdahulu', '29'), ('2.9 Sintesis Research Gap dan Posisi Penelitian', '32'),
        ('BAB III OBJEK DAN METODOLOGI PENELITIAN', '34'), ('3.1 Objek Penelitian', '34'), ('3.2 Metodologi Penelitian', '36'),
        ('3.3 Tahapan dan Rancangan Penelitian', '37'), ('3.4 Rancangan Artefak dan Arsitektur Viero Agent', '38'),
        ('3.5 Mekanisme Keputusan, Execution, dan Safety', '42'), ('3.5.1 Fresh State Confirmation dan Authorization Boundary', '42'),
        ('3.5.2 Alur Keputusan Deterministik', '43'), ('3.5.3 Screening dan Verifikasi Pool', '45'), ('3.5.4 Position Planning dan Execution', '46'),
        ('3.5.5 Position Lifecycle dan Recovery', '47'), ('3.5.6 Safety dan Operator Boundary', '48'), ('3.6 Pengujian Perangkat Lunak dan Evaluasi Artefak', '49'),
        ('Daftar Pustaka', '55')]
    add_static_index(doc, 'Daftar Isi', toc_entries)
    doc.add_page_break()
    add_static_index(doc, 'Daftar Gambar', [('Gambar 3.1 Alur Penelitian Rancang Bangun dan Evaluasi Artefak Viero Agent', '38'), ('Gambar 3.2 Arsitektur Konseptual Viero Agent dan Aliran Data/Execution', '42'), ('Gambar 3.3 Alur Keputusan, Execution, dan Recovery Viero Agent', '45')])
    doc.add_page_break()
    add_static_index(doc, 'Daftar Tabel', [('Tabel 1.1 Jadwal Penelitian', '9'), ('Tabel 2.1 Ringkasan Penelitian Terdahulu', '30'), ('Tabel 3.1 Jaringan dan Protokol yang Didukung Viero', '34'), ('Tabel 3.2 Komponen dan Tanggung Jawab Arsitektur Viero Agent', '39'), ('Tabel 3.3 Pipeline Keputusan dan Execution Viero Agent', '44'), ('Tabel 3.4 Lapisan Pemeriksaan Safety', '48'), ('Tabel 3.5 Level Pengujian dan Evaluasi Artefak', '49'), ('Tabel 3.6 Matriks Skenario dan Evaluasi E1–E12', '52')])
    # Main content section with Arabic numbering
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.footer.is_linked_to_previous = False
    new_sec.different_first_page_header_footer = False
    new_sec.page_width = Cm(21); new_sec.page_height = Cm(29.7)
    new_sec.left_margin = Cm(4); new_sec.right_margin = Cm(3); new_sec.top_margin = Cm(3); new_sec.bottom_margin = Cm(3)
    add_page_number_footer(new_sec, roman=False)
    add_content(doc)
    add_bibliography_format(doc)
    # Set main header/footer and update fields.
    settings = doc.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
    doc.core_properties.title = TITLE
    doc.core_properties.author = 'Faisal Hidayat Sukma'
    doc.core_properties.subject = 'Proposal Tugas Akhir'
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__': build()
